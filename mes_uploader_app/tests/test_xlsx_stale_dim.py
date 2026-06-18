# -*- coding: utf-8 -*-
"""Regression: file .xlsx có <dimension> KHAI THIẾU dòng (máy đo append nhưng
không cập nhật dimension) -> vẫn phải đọc ĐÚNG dòng CUỐI (mới nhất).

Trước đây dùng openpyxl read_only (tin theo <dimension>) nên BỎ SÓT dòng mới ->
đọc nhầm dòng cũ (vd lần 2 thực tế OK nhưng lấy lại NG của lần 1). Nay đọc đầy
đủ sheetData (read_only=False).

Chạy:  python -m tests.test_xlsx_stale_dim
"""

import io
import os
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from mes_uploader import data_reader, excel_export as xe

HDR = ["Time", "Judge", "IspTime", "Data01", "Data02"]


def _stale_dim_xlsx(path, rows, fake_ref="A1:Z2"):
    """Tạo .xlsx (header + rows) rồi SỬA <dimension> khai thiếu (mô phỏng máy đo)."""
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(HDR)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); wb.close()
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for name in zin.namelist():
        content = zin.read(name)
        if name == "xl/worksheets/sheet1.xml":
            text = content.decode("utf-8")
            text = re.sub(r'(<dimension ref=")[^"]*(")', r"\g<1>%s\g<2>" % fake_ref, text)
            content = text.encode("utf-8")
        zout.writestr(name, content)
    zout.close(); zin.close()
    with open(path, "wb") as f:
        f.write(out.getvalue())


def main():
    tmp = tempfile.mkdtemp(prefix="stale_dim_")
    p = os.path.join(tmp, "CCD1_NearStack.xlsx")
    # 2 dòng dữ liệu: lần 1 NG (10,11), lần 2 OK (20,21); dimension khai chỉ tới
    # dòng 2 (tức 1 data row) — mô phỏng máy ghi thêm dòng mà không sửa dimension.
    _stale_dim_xlsx(p, [["09:10", "NG", 500, 10.0, 11.0],
                        ["09:11", "OK", 500, 20.0, 21.0]])

    # sanity: read_only=True (CÁCH CŨ) sẽ BỎ SÓT dòng mới -> đọc nhầm NG
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ro = [r for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    ro_data = [r for r in ro[1:] if r and any(c not in (None, "") for c in r)]
    print("== mô phỏng dimension sai ==")
    print("  read_only=True thấy %d data row (lỗi cũ: chỉ thấy dòng đầu)" % len(ro_data))

    # 1) data_reader đọc ĐÚNG dòng CUỐI (mới): OK (20,21), KHÔNG phải NG cũ
    r = data_reader.read_latest_measurement(p)
    assert r["judge"] == "OK", r["judge"]
    assert r["values"] == [20.0, 21.0], r["values"]
    print("  data_reader -> judge=OK values=[20.0,21.0]  ✔ (đọc đúng dòng mới)")

    # 2) excel_export.extract_source cũng lấy dòng CUỐI (mới)
    hdr, last = xe.extract_source(open(p, "rb").read())
    vals = [c["value"] for c in last]
    assert vals[1] == "OK" and vals[3] == 20.0 and vals[4] == 21.0, vals
    print("  extract_source -> dòng cuối OK (20,21)  ✔")

    print("\nTEST XLSX-STALE-DIM PASS ✔")


if __name__ == "__main__":
    main()
