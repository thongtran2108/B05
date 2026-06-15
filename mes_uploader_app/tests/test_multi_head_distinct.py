# -*- coding: utf-8 -*-
"""SN nhiều đầu: mỗi đầu lưu ĐÚNG measurement/ảnh của LẦN ĐO ĐÓ (không trùng).

Kịch bản đúng như yêu cầu: 2 đầu 8X, giữa 2 lần tín hiệu PLC máy đo ghi THÊM
1 dòng + 1 ảnh cho mỗi đầu:
  - đầu #1 đo NG  -> lưu dòng/ảnh NG
  - đầu #2 đo OK  -> lưu dòng/ảnh OK
Kỳ vọng:
  - data đọc cho 2 đầu KHÁC nhau (NG rồi OK), không lấy trùng dòng cuối.
  - Excel có 2 dòng KHÁC nhau (NG, OK) — mỗi đầu 1 dòng của mình.
  - Ảnh: đầu #1 -> _Left_Failed_#1, đầu #2 -> _Left_Passed_#2 (mỗi đầu ảnh riêng).
  - result POST lên MES = FAIL (1 đầu NG là FAIL).

Chạy:  python -m tests.test_multi_head_distinct
"""

import datetime
import glob
import os
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from PIL import Image

from mes_uploader import mes_api
from mes_uploader.config import AppConfig, MaterialConfig, PathConfig
from mes_uploader.core.side_worker import SideWorker
from mes_uploader.hardware.plc_client import MockPlcClient

HEADER = "Time,Judge,IspTime,Data01,Data02\n"


def _append_row(path, time_s, judge, v1, v2):
    with open(path, "a", encoding="utf-8") as f:
        f.write("%s,%s,500,%s,%s\n" % (time_s, judge, v1, v2))


def _mkjpg(path, size):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    Image.new("RGB", size, (20, 90, 160)).save(path, "JPEG")


def main():
    root = tempfile.mkdtemp(prefix="multi_head_")
    day = datetime.datetime.now().strftime("%Y%m%d")

    # File đo gốc của bên TRÁI (CCD1), đầu 8X: header + 1 dòng CŨ (mốc = 1).
    data_dir = os.path.join(root, "8X", "data", day)
    os.makedirs(data_dir, exist_ok=True)
    src_csv = os.path.join(data_dir, "CCD1_test.csv")
    with open(src_csv, "w", encoding="utf-8") as f:
        f.write(HEADER)
    _append_row(src_csv, "09:00:00", "OK", 1.0, 1.0)     # dòng CŨ (SN trước)

    img_src = os.path.join(root, "src8x")
    img_up = os.path.join(root, "up8x")
    xl_dir = os.path.join(root, "xl8x")

    cfg = AppConfig()
    cfg.simulation = True
    cfg.poll_interval_ms = 20
    cfg.paths = PathConfig(base_dir=root)
    cfg.paths.require_today = False
    cfg.materials = [MaterialConfig("ABC", heads_8x=2)]
    cfg.api.api_8x.url = "http://mes/8x/upload"
    cfg.images.enabled = True
    cfg.images.img_8x.source_dir = img_src
    cfg.images.img_8x.upload_dir = img_up
    cfg.excel.enabled = True
    cfg.excel.output_dir_8x = xl_dir

    payloads = []
    mes_api.post_payload = lambda url, payload, **kw: (payloads.append(payload),
                                                       (True, 200, "OK"))[1]

    readings = []

    def on_event(et, **d):
        if et == "reading":
            readings.append(d)

    w = SideWorker("left", cfg, MockPlcClient(), on_event)
    w.start()
    w.arm(cfg.materials[0], "8X")        # 2 đầu 8X
    time.sleep(0.1)
    w.submit_sn("SN-MH-9")               # CHỐT mốc tại đây: row_base = 1
    time.sleep(0.2)

    # --- đầu #1: máy đo ghi thêm dòng NG + ảnh NG, rồi báo tín hiệu ---
    _append_row(src_csv, "09:10:00", "NG", 10.0, 11.0)
    _mkjpg(os.path.join(img_src, "Image", day, "CCD1", "NG", "h1.jpg"), (11, 11))
    w.simulate_trigger()
    time.sleep(0.4)

    # --- đầu #2: máy đo ghi thêm dòng OK + ảnh OK, rồi báo tín hiệu ---
    _append_row(src_csv, "09:11:00", "OK", 20.0, 21.0)
    _mkjpg(os.path.join(img_src, "Image", day, "CCD1", "OK", "h2.jpg"), (22, 22))
    w.simulate_trigger()
    time.sleep(0.6)                      # chờ luồng nền (ảnh) xử lý xong TRƯỚC stop
    w.stop()

    # 1) DATA mỗi đầu KHÁC nhau: đầu #1 = NG (10,11), đầu #2 = OK (20,21)
    print("== data theo từng đầu ==")
    r = {d["index"]: d for d in readings}
    assert set(r) == {1, 2}, "phải có 2 đầu, nhận: %r" % sorted(r)
    assert r[1]["judge"] == "NG" and list(r[1]["values"]) == [10.0, 11.0], r[1]
    assert r[2]["judge"] == "OK" and list(r[2]["values"]) == [20.0, 21.0], r[2]
    print("  đầu#1 NG (10,11) | đầu#2 OK (20,21)  ✔")

    # 2) result POST = FAIL (1 đầu NG là FAIL)
    print("\n== result POST lên MES ==")
    assert payloads, "phải có POST"
    assert payloads[-1]["result"] == "FAIL", payloads[-1]
    # timer 2 đầu khác nhau (không trùng)
    timer = payloads[-1]["timer"]
    assert "10.0-11.0" in timer and "20.0-21.0" in timer, timer
    print("  result=FAIL | timer có cả (10.0-11.0) và (20.0-21.0)  ✔")

    # 3) EXCEL: 2 dòng KHÁC nhau (NG rồi OK), tên file CCD1->Left
    print("\n== Excel mỗi đầu 1 dòng của mình ==")
    xls = glob.glob(os.path.join(xl_dir, "**", "*.xlsx"), recursive=True)
    assert xls, "phải tạo file Excel"
    assert os.path.basename(xls[0]).startswith("Left_"), xls[0]
    wb = openpyxl.load_workbook(xls[0])
    ws = wb.active
    try:
        rows = [(ws.cell(rr, 1).value, ws.cell(rr, 3).value)   # (SN, Judge)
                for rr in range(2, ws.max_row + 1) if ws.cell(rr, 1).value]
        assert len(rows) == 2, "2 đầu -> 2 dòng, nhận %d" % len(rows)
        assert all(sn == "SN-MH-9" for sn, _ in rows), rows
        assert [j for _, j in rows] == ["NG", "OK"], rows   # đúng thứ tự + KHÁC nhau
        print("  2 dòng: NG rồi OK (file %s)  ✔" % os.path.basename(xls[0]))
    finally:
        wb.close()

    # 4) ẢNH: đầu#1 -> _Left_Failed_#1 (NG), đầu#2 -> _Left_Passed_#2 (OK)
    print("\n== ảnh mỗi đầu riêng ==")
    names = sorted(fn for _d, _s, fns in os.walk(img_up) for fn in fns)
    assert len(names) == 2, "2 đầu -> 2 ảnh, nhận: %r" % names
    fail1 = [n for n in names if "_Left_Failed_#1" in n]
    pass2 = [n for n in names if "_Left_Passed_#2" in n]
    assert fail1 and pass2, names
    # nội dung đúng ảnh của từng đầu: #1 = ảnh NG (11x11), #2 = ảnh OK (22x22)
    with Image.open(os.path.join(img_up, day, fail1[0])) as im:
        assert im.size == (11, 11), im.size
    with Image.open(os.path.join(img_up, day, pass2[0])) as im:
        assert im.size == (22, 22), im.size
    print("  #1=%s (11x11) | #2=%s (22x22)  ✔" % (fail1[0], pass2[0]))

    print("\nTEST MULTI-HEAD-DISTINCT PASS ✔")


if __name__ == "__main__":
    main()
